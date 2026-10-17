"""Extract event metadata, menu and chef's raw-material demand from RESHEEN event order workbooks.

Usage: python extract_events.py <folder_with_xlsx> <out_dir>
Outputs: events_raw.csv (one row per file), demand_long.csv (event x raw item demand), menu_long.csv (event x dish)
"""
import csv, glob, os, re, sys, warnings, datetime as dt
import openpyxl

warnings.filterwarnings("ignore")

SECTION_HEADERS = {
    "WELCOME DRINK": "welcome", "BITES": "bites", "SALAD": "salad", "DRESSING": "dressing",
    "SOUP": "soup", "HOT CHAFFER": "hot_chaffer", "CHICKEN": "chicken", "PORK": "pork",
    "SHELL FISH": "shellfish", "FISH": "fish", "CUTTLEFISH": "cuttlefish", "MUTTON": "mutton",
    "BEEF": "beef", "VEGITABLE": "vegetable", "VEGETABLE": "vegetable", "ACTION": "action",
    "CONDIMENTS": "condiments", "DESSERTS": "desserts", "SPECIAL": "special",
}


def parse_filename(name):
    base = os.path.splitext(name)[0]
    m = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", base)
    date = None
    if m:
        d, mo, y = map(int, m.groups())
        if y == 2626:  # typo in one filename
            y = 2026
        date = dt.date(y, mo, d)
    hall = "RG" if re.search(r"\bRG\b|RG\)|RG-|RG ", base) else "BH"
    low = base.lower()
    session = "night" if "night" in low else "day"
    return date, hall, session


def norm(s):
    return re.sub(r"\s+", " ", str(s)).strip()


def parse_menu(ws):
    info = {"pax_booked": ws["E9"].value, "event_no": ws["E10"].value, "date_in_sheet": ws["E8"].value}
    typ = ws["F8"].value or ""
    info["event_type_raw"] = norm(typ.split(":", 1)[-1]) if ":" in str(typ) else norm(typ)
    loc = ws["F9"].value or ""
    info["client_location"] = norm(str(loc).split(":", 1)[-1])
    actual = None
    dishes = []  # (section, dish, qty, unit)
    for col in "CEG":
        section = None
        for r in range(11, 101):
            v = ws[f"{col}{r}"].value
            if v is None or str(v).strip() == "":
                continue
            s = norm(v)
            ma = re.search(r"actual\s*count\s*-?\s*(\d+)", s, re.I)
            if ma:
                actual = int(ma.group(1))
                continue
            if s.upper() in SECTION_HEADERS:
                section = SECTION_HEADERS[s.upper()]
                continue
            if section is None:
                continue
            nxt = chr(ord(col) + 1)
            q = ws[f"{nxt}{r}"].value
            unit = ws[f"{chr(ord(col) + 2)}{r}"].value
            q = q if isinstance(q, (int, float)) else None
            dishes.append((section, s, q, unit if isinstance(unit, str) else None))
    # Actual count can also sit in other columns of the header area
    if actual is None:
        for row in ws.iter_rows(min_row=8, max_row=20, max_col=12, values_only=True):
            for v in row:
                if isinstance(v, str):
                    ma = re.search(r"actual\s*count\s*-?\s*(\d+)", v, re.I)
                    if ma:
                        actual = int(ma.group(1))
    info["pax_actual"] = actual
    return info, dishes


def parse_stock(ws):
    rate = ws["O10"].value
    items = []
    sub = None
    for r in range(13, ws.max_row + 1):
        c = ws[f"C{r}"].value
        if isinstance(c, str) and "Sub Category" in c:
            sub = norm(c.split(":", 1)[-1])
            continue
        code, desc, unit, demand = (ws[f"F{r}"].value, ws[f"G{r}"].value, ws[f"J{r}"].value, ws[f"K{r}"].value)
        if desc and isinstance(demand, (int, float)) and demand > 0:
            items.append((sub, norm(code) if code else "", norm(desc), norm(unit) if unit else "", float(demand)))
    return rate, items


def main(src, out):
    os.makedirs(out, exist_ok=True)
    ev_rows, dem_rows, menu_rows = [], [], []
    for path in sorted(glob.glob(os.path.join(src, "*.xlsx"))):
        name = os.path.basename(path)
        wb = openpyxl.load_workbook(path, data_only=True)
        date, hall, session = parse_filename(name)
        info, dishes = parse_menu(wb["Party Menu"])
        rate, items = parse_stock(wb["StockBalance"])
        eid = f"E{len(ev_rows) + 1:03d}"
        sec_counts = {}
        for sec, *_ in dishes:
            sec_counts[sec] = sec_counts.get(sec, 0) + 1
        row = {"event_id": eid, "file": name, "date": date, "hall": hall, "session": session,
               "rate_lkr": rate, **info}
        for sec in sorted(set(SECTION_HEADERS.values())):
            row[f"n_{sec}"] = sec_counts.get(sec, 0)
        ev_rows.append(row)
        for sec, dish, q, unit in dishes:
            menu_rows.append({"event_id": eid, "section": sec, "dish": dish, "qty": q, "unit": unit})
        for sub, code, desc, unit, demand in items:
            dem_rows.append({"event_id": eid, "sub_category": sub, "item_code": code, "item": desc,
                             "unit": unit, "demand": round(demand, 4)})
    for fname, rows in (("events_raw.csv", ev_rows), ("demand_long.csv", dem_rows), ("menu_long.csv", menu_rows)):
        with open(os.path.join(out, fname), "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
            w.writeheader()
            w.writerows(rows)
        print(fname, len(rows))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
