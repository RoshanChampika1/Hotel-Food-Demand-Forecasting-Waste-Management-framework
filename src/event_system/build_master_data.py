"""Build the event-order system's master data from the hotel's own event-order workbooks.

Reads every data/raw/*.xlsx and writes data/master/ (git-ignored, hotel data):
  items.csv        item code, name, unit, stock group (Fresh / Dry / Meat), StockBalance order
  recipes.csv      dish, section, basis, item code, qty (per 100 guests, or per kg for bites)
  dishes.csv       every dish seen on a menu or in a recipe book, its menu section and how often it was served
  bites.csv        typical kg divisor for each bite (kg = guests / divisor), from the chef's formulas
  past_menus.json  the printed menu of each past event (to copy as a starting point)

Recipes come from the newest workbook first ("New Recepies" > "Recepies" > "RG-Recepies"),
so the latest version of a recipe wins. Recipes are linked to menu sections through the hidden
Party Menu matrix (columns S to PC), exactly as the workbook's SUMIFS formulas do.

Usage: python src/event_system/build_master_data.py data/raw data/master
"""
import json
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime

import openpyxl
import pandas as pd

# Printed Party Menu headers -> menu section key (the kitchen's own words)
HEADER_MAP = {
    "WELCOME DRINK": "welcome", "WELCOME DRINKS": "welcome", "BITES": "bites", "SPECIAL": "special",
    "SALAD": "salad", "SALADS": "salad", "DRESSING": "dressing", "SOUP": "soup",
    "HOT CHAFFER": "hot_chaffer", "CHICKEN": "chicken", "PORK": "pork", "FISH": "fish",
    "SHELL FISH": "fish", "SHELLFISH": "fish", "CUTTLEFISH": "fish", "MUTTON": "mutton", "BEEF": "beef",
    "VEGITABLE": "vegetable", "VEGETABLE": "vegetable", "ACTION": "action",
    "CONDIMENTS": "condiments", "DESSERTS": "desserts", "DESSERT": "desserts",
}
# Party Menu matrix labels (column S) -> menu section key
MATRIX_MAP = {
    "welcome": "welcome", "veg.salads": "salad", "nonveg.salads": "salad", "mirrorsalads": "salad",
    "salads": "salad", "dressing": "dressing", "hot chaffers": "hot_chaffer", "chicken": "chicken",
    "pork": "pork", "fish": "fish", "vegetable": "vegetable", "action": "action",
    "condimants": "condiments", "condiments": "condiments", "desserts": "desserts", "soup": "soup",
    "bites": "bites", "additional": "special",
}
MEAT_WORDS = re.compile(
    r"chicken|pork|fish|prawn|cuttle|calamari|squid|sausage|bacon|ham\b|salami|mutton|beef|lamb|tuna|"
    r"thalapath|linna|salaya|seer|crab|mussel|lobster|meat|brockwest|drumst|thigh|wings|liver|bird",
    re.I)
NOT_MEAT = re.compile(r"powder|seasoning|cube|stock|sauce|sause|paste|masala|maldive|moldive|dry|cracker|soya|"
                      r"cand |crushed", re.I)
GENERIC_NAMES = {"dressing", "dressings"}


def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def key(s):
    return norm(s).lower().replace("’", "'")


def file_date(name):
    m = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", name)
    if not m:
        return datetime(2000, 1, 1)
    d, mth, y = map(int, m.groups())
    y = 2026 if y == 2626 else y  # typo in one file name
    return datetime(y, mth, d)


def formula_text(v):
    return v.text if hasattr(v, "text") else v


def read_items(wb, items):
    """Item code, name, unit and Dry/Fresh group from the StockBalance sheet, in sheet order."""
    ws = wb["StockBalance"]
    group = None
    for r in range(13, ws.max_row + 1):
        c = norm(ws.cell(r, 3).value)
        if "Sub Category" in c:
            group = "Dry" if "Dry" in c else "Fresh" if "Fresh" in c else c.split(":")[-1].strip()
        code, name, unit = norm(ws.cell(r, 6).value), norm(ws.cell(r, 7).value), norm(ws.cell(r, 10).value)
        if not code or not name or code == "Item Code":
            continue
        k = code.upper().replace(" ", "")
        if k not in items:
            items[k] = dict(code=code, name=name, unit=unit, sheet_group=group, order=len(items))
        elif unit and not items[k]["unit"]:
            items[k]["unit"] = unit


def matrix_blocks(wb):
    """Parse the hidden Party Menu matrix: which recipe rows belong to which section, and the basis."""
    ws = wb["Party Menu"]
    blocks, label = [], None
    for r in range(11, 200):
        s = ws.cell(r, 19).value
        if s:
            label = norm(s).lower()
        f = formula_text(ws.cell(r, 22).value)
        if not isinstance(f, str) or "SUMIFS" not in f:
            continue
        m = re.search(r"'?([A-Za-z\- ]*Recepies)'?!\$D\$(\d+)(?::\$D\$(\d+))?", f)
        if not m:
            continue
        sheet, r1 = m.group(1), int(m.group(2))
        r2 = int(m.group(3) or r1)
        tail = f.split("))")[-1]
        if re.search(r"/100\*\$H", tail):
            basis = "per_100_count"   # e.g. Ice Coffee: per 100 glasses
        elif re.search(r"\*\$H", tail):
            basis = "per_unit_count"  # e.g. Milk rice: per kg ordered
        elif re.search(r"/100\*\$F", tail):
            basis = "per_100_count_bite"
        elif re.search(r"\*\$F", tail):
            basis = "per_kg_bite"     # bites: recipe is per 1 kg, kg = guests / divisor
        else:
            basis = "per_100_pax"
        section = MATRIX_MAP.get(label, "special") if label else "special"
        if basis.endswith("bite"):
            section = "bites"
        blocks.append(dict(sheet=sheet, r1=r1, r2=r2, basis=basis, section=section,
                           label=norm(ws.cell(r, 21).value)))
    return blocks


def read_recipes(wb, blocks):
    """Return {(dish_key, section): recipe dict} for this workbook."""
    out = {}
    for b in blocks:
        if b["sheet"] not in wb.sheetnames:
            continue
        ws = wb[b["sheet"]]
        for r in range(b["r1"], b["r2"] + 1):
            dish, code, item, qty, unit = (ws.cell(r, c).value for c in range(1, 6))
            if not dish or not code or not isinstance(qty, (int, float)):
                continue
            if key(dish) in GENERIC_NAMES:
                continue  # rows named just "Dressing" never match a menu dish in the workbook, so they add nothing
            k = (key(dish), b["section"])
            rec = out.setdefault(k, dict(dish=norm(dish), section=b["section"], basis=b["basis"],
                                         sheet=b["sheet"], lines={}))
            ck = norm(code).upper().replace(" ", "")
            if ck in rec["lines"]:  # same item twice in one recipe: add up, as SUMIFS does
                rec["lines"][ck]["qty"] += float(qty)
            else:
                rec["lines"][ck] = dict(item_code=norm(code), item_name=norm(item), qty=float(qty), unit=norm(unit))
    return out


def read_menu(wb):
    """The printed menu: header cells in columns C, E, G; dishes below each header."""
    ws = wb["Party Menu"]
    sections = defaultdict(list)
    bites = {}
    for col in (3, 5, 7):
        current = None
        for r in range(12, 92):
            v = norm(ws.cell(r, col).value)
            if not v:
                continue
            up = v.upper()
            if up in HEADER_MAP:
                current = HEADER_MAP[up]
                continue
            if current is None or v.startswith("=") or up.startswith("ACTUAL COUNT") or "SERVING" in up:
                continue
            sections[current].append(v)
            if current == "bites":
                f = formula_text(ws.cell(r, col + 1).value)
                m = re.search(r"E9\s*/\s*([\d.]+)", str(f or ""))
                if m:
                    bites[v] = float(m.group(1))
    head = {}
    head["pax"] = ws["E9"].value
    t = norm(ws["F8"].value)
    head["event_type"] = t.split(":", 1)[-1].strip() if ":" in t else t
    loc = norm(ws["F9"].value)
    head["location"] = loc.split(":", 1)[-1].strip() if ":" in loc else loc
    srv = ""
    for c in ("F10", "G10", "H10"):
        if "serv" in norm(ws[c].value).lower():
            srv = re.sub(r"(?i)food\s*serv\w*\s*(time)?\s*-?\s*", "", norm(ws[c].value))
    head["serving_time"] = srv
    return dict(sections), bites, head


def stock_group(name, sheet_group):
    if MEAT_WORDS.search(name) and not NOT_MEAT.search(name):
        return "Meat"
    return "Dry" if sheet_group == "Dry" else "Fresh"


def main(raw_dir, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    files = sorted((f for f in os.listdir(raw_dir)
                    if f.lower().endswith(".xlsx") and f != "Event Order New imo.xlsx"),  # test file
                   key=file_date, reverse=True)
    items, recipes, menus = {}, {}, []
    bite_div, dish_count, dish_last, dish_section = defaultdict(list), Counter(), {}, {}
    for f in files:
        wb = openpyxl.load_workbook(os.path.join(raw_dir, f))
        read_items(wb, items)
        for k, rec in read_recipes(wb, matrix_blocks(wb)).items():
            recipes.setdefault(k, rec | {"source": f})  # newest file wins
        sections, bites, head = read_menu(wb)
        for d, div in bites.items():
            bite_div[key(d)].append(div)
        hall = "RG" if "RG" in f else "BH"
        session = "night" if "night" in f.lower() else "day"
        menus.append(dict(file=f, date=file_date(f).strftime("%Y-%m-%d"), hall=hall, session=session,
                          **head, sections=sections))
        for sec, ds in sections.items():
            for d in ds:
                dish_count[(key(d), sec)] += 1
                dish_last.setdefault((key(d), sec), file_date(f).strftime("%Y-%m-%d"))
                dish_section.setdefault((key(d), sec), d)
        print(f"read {f}")

    # items
    rows = []
    for k, it in items.items():
        rows.append(dict(item_key=k, item_code=it["code"], item_name=it["name"], unit=it["unit"] or "Kg",
                         group=stock_group(it["name"], it["sheet_group"]), order=it["order"]))
    # recipe items missing from StockBalance
    known = set(items)
    for rec in recipes.values():
        for ck, ln in rec["lines"].items():
            if ck not in known:
                known.add(ck)
                grp = "Dry" if ck.startswith("D") else "Fresh"
                rows.append(dict(item_key=ck, item_code=ln["item_code"], item_name=ln["item_name"],
                                 unit=ln["unit"] or "Kg", group=stock_group(ln["item_name"], grp), order=9000 + len(rows)))
    items_df = pd.DataFrame(rows)
    items_df.to_csv(os.path.join(out_dir, "items.csv"), index=False)

    # recipes
    rrows = []
    for (dk, sec), rec in recipes.items():
        for ck, ln in rec["lines"].items():
            rrows.append(dict(dish_key=dk, dish=rec["dish"], section=sec, basis=rec["basis"], item_key=ck,
                              item_code=ln["item_code"], item_name=ln["item_name"], qty=ln["qty"],
                              recipe_unit=ln["unit"], source=rec["source"], sheet=rec["sheet"]))
    pd.DataFrame(rrows).to_csv(os.path.join(out_dir, "recipes.csv"), index=False)

    # dishes: union of menu dishes and recipe dishes
    drows = {}
    for (dk, sec), rec in recipes.items():
        drows[(dk, sec)] = dict(dish_key=dk, dish=rec["dish"], section=sec, basis=rec["basis"], has_recipe=1)
    for (dk, sec), name in dish_section.items():
        d = drows.setdefault((dk, sec), dict(dish_key=dk, dish=name, section=sec, basis="per_100_pax", has_recipe=0))
        if not d["has_recipe"]:
            # a menu dish may have its recipe under another section (e.g. a vegetable dish listed as a condiment)
            other = [r for (k2, s2), r in recipes.items() if k2 == dk]
            if other:
                d.update(has_recipe=1, basis=other[0]["basis"], recipe_section=other[0]["section"])
    dishes = pd.DataFrame(drows.values())
    dishes["times_served"] = [dish_count.get((r.dish_key, r.section), 0) for r in dishes.itertuples()]
    dishes["last_served"] = [dish_last.get((r.dish_key, r.section), "") for r in dishes.itertuples()]
    if "recipe_section" not in dishes:
        dishes["recipe_section"] = None
    dishes["recipe_section"] = dishes.recipe_section.fillna(dishes.section)
    dishes.sort_values(["section", "times_served", "dish"], ascending=[True, False, True]).to_csv(
        os.path.join(out_dir, "dishes.csv"), index=False)

    # bites divisors (most common value the chef used)
    brows = [dict(dish_key=k, divisor=Counter(v).most_common(1)[0][0], n=len(v)) for k, v in bite_div.items()]
    pd.DataFrame(brows).to_csv(os.path.join(out_dir, "bites.csv"), index=False)

    with open(os.path.join(out_dir, "past_menus.json"), "w") as fh:
        json.dump(menus, fh, indent=1, default=str)
    print(f"\n{len(files)} workbooks -> {len(items_df)} items, {len(recipes)} recipes, "
          f"{len(dishes)} dishes, {len(brows)} bites, {len(menus)} past menus -> {out_dir}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "data/raw", sys.argv[2] if len(sys.argv) > 2 else "data/master")
