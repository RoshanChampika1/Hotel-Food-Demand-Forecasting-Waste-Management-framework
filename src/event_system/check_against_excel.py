"""Check the system's StockBalance against the demand Excel calculated for a past event.

Takes the event's own menu and booked pax, computes the demand with the system and compares it with the
cached values in the workbook's StockBalance (column K). Run after build_master_data.py.

Usage: python src/event_system/check_against_excel.py "data/raw/Event Order (10.10.2026 - BH Night).xlsx"
"""
import sys

import openpyxl
import pandas as pd

from order_calc import bite_divisor, event_demand, load_master


def main(path):
    master = load_master()
    name = path.replace("\\", "/").split("/")[-1]
    m = next(p for p in master["past_menus"] if p["file"] == name)
    pax = int(m["pax"])
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["StockBalance"]
    excel = {}
    for r in range(15, ws.max_row + 1):
        code, k = ws.cell(r, 6).value, ws.cell(r, 11).value
        if code and isinstance(k, (int, float)) and k > 0:
            excel[str(code).upper().replace(" ", "")] = float(k)
    pmv = wb["Party Menu"]
    bite_kg = {}
    for r in range(17, 24):
        d, kg = pmv.cell(r, 5).value, pmv.cell(r, 6).value
        if d and isinstance(kg, (int, float)):
            bite_kg[str(d).strip()] = float(kg)
    menu = {s: [d for d in ds] for s, ds in m["sections"].items()}
    counts = {}
    for r in range(60, 80):
        d, n = pmv.cell(r, 7).value, pmv.cell(r, 8).value
        if d and isinstance(n, (int, float)):
            counts[str(d).strip()] = n
    _, sb, missing = event_demand(master, menu, pax, bite_kg, counts)
    ours = dict(zip(sb.item_key, sb.demand))
    keys = sorted(set(excel) | set(ours))
    df = pd.DataFrame({"item": keys, "excel": [excel.get(k, 0) for k in keys], "system": [ours.get(k, 0) for k in keys]})
    df["name"] = df.item.map(dict(zip(sb.item_key, sb.description)))
    df["diff"] = df.system - df.excel
    ok = (df["diff"].abs() <= 0.01 + 0.01 * df.excel).sum()
    print(f"{name}: {pax} pax, {len(df)} items, {ok} match (within 1%), dishes without recipe: {missing}")
    print(df[(df["diff"].abs() > 0.01 + 0.01 * df.excel)].round(3).to_string(index=False))


if __name__ == "__main__":
    main(sys.argv[1])
